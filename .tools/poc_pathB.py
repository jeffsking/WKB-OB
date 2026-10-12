# -*- coding: utf-8 -*-
"""路径 B 端到端验证：get_media_info -> 下载原文件 -> 本地解析。
对比 MCP fetch_media_content 的解析质量，判断脚本化批量提取是否可行。
"""
import json, os, subprocess, urllib.request, sys, time

SKILL_DIR = r"C:\Users\Chris Wang\.workbuddy\skills\ima-skill"
NODE = r"C:\Users\Chris Wang\.workbuddy\binaries\node\versions\22.22.2-3\node.exe"
HOME = os.path.expanduser("~")

def creds():
    cid = open(os.path.join(HOME, ".config", "ima", "client_id"), encoding="utf-8").read().strip()
    key = open(os.path.join(HOME, ".config", "ima", "api_key"), encoding="utf-8").read().strip()
    return cid, key

def get_media_info(media_id):
    cid, key = creds()
    opts = json.dumps({"clientId": cid, "apiKey": key})
    p = subprocess.run([NODE, os.path.join(SKILL_DIR, "ima_api.cjs"),
                        "openapi/wiki/v1/get_media_info",
                        json.dumps({"media_id": media_id}), opts],
                       capture_output=True, text=True, encoding="utf-8")
    return json.loads(p.stdout)

def download(url, headers, out_path):
    req = urllib.request.Request(url)
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    with open(out_path, "wb") as f:
        f.write(data)
    return len(data)

def parse_pdf(path):
    from PyPDF2 import PdfReader
    r = PdfReader(path)
    pages = [(i, (pg.extract_text() or "")) for i, pg in enumerate(r.pages)]
    return pages

os.makedirs(".tools/_poc_dl", exist_ok=True)

CASES = [
    ("pdf",   "pdf_3867e2b60070d7dc3f413086d0b0036e_baf5396e223f6aacaea1aab9cadc79af001a693ee6c0644e", "pdf"),
    ("ppt",   "ppt_3867e2b60070d7dc3f413086d0b0036e_76e413aa1f6d3d700d5f14570614d4a2001a693ee6c0644e", "pptx"),
    ("word",  "word_3867e2b60070d7dc3f413086d0b0036e_d2d05fb7e51b6b61e9e3061be410edf8001a693ee6c0644e", "docx"),
    ("excel", "excel_3867e2b60070d7dc3f413086d0b0036e_ac1ec2cf271de520ce26bac558cadc3d001a693ee6c0644e", "xlsx"),
]

print("=" * 70)
print("路径 B 端到端实测：REST get_media_info -> 下载 -> 本地解析")
print("=" * 70)

for kind, media_id, ext in CASES:
    t0 = time.time()
    try:
        info = get_media_info(media_id)
    except Exception as e:
        print(f"[{kind}] get_media_info 失败: {e}")
        continue
    if info.get("code") != 0:
        print(f"[{kind}] code={info.get('code')} msg={info.get('msg')}")
        continue
    ui = info["data"].get("url_info") or {}
    url, hdrs = ui.get("url"), ui.get("headers")
    if not url:
        print(f"[{kind}] url_info 为空（media_type={info['data'].get('media_type')}）")
        continue
    out = os.path.join(".tools/_poc_dl", f"{kind}.{ext}")
    try:
        n = download(url, hdrs, out)
    except Exception as e:
        print(f"[{kind}] 下载失败: {e}")
        continue
    dt = time.time() - t0
    print(f"\n[{kind}] 下载 {n:,} B -> {out}  (耗时 {dt:.1f}s)")

    # 本地解析
    try:
        if ext == "pdf":
            pages = parse_pdf(out)
            total = sum(len(t) for _, t in pages)
            print(f"  解析: {len(pages)} 页, 提取正文 {total:,} 字符")
            if pages:
                head = pages[0][1][:200].replace("\n", " / ")
                print(f"  首页前 200 字: {head}")
        elif ext == "pptx":
            from pptx import Presentation
            pr = Presentation(out)
            txt = []
            for s in pr.slides:
                for sh in s.shapes:
                    if sh.has_text_frame:
                        txt.append(sh.text_frame.text)
            joined = "\n".join(txt)
            print(f"  解析: {len(pr.slides)} 页, 提取正文 {len(joined):,} 字符")
            print(f"  前 200 字: {joined[:200]}".replace("\n", " / "))
        elif ext == "docx":
            import docx
            d = docx.Document(out)
            txt = "\n".join(p.text for p in d.paragraphs)
            print(f"  解析: {len(d.paragraphs)} 段落, 提取正文 {len(txt):,} 字符")
            print(f"  前 200 字: {txt[:200]}".replace("\n", " / "))
        elif ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(out, data_only=True)
            print(f"  解析: {len(wb.sheetnames)} 个 sheet: {wb.sheetnames}")
            ws = wb[wb.sheetnames[0]]
            print(f"  首个 sheet 尺寸: {ws.max_row} 行 x {ws.max_column} 列")
    except Exception as e:
        print(f"  解析失败: {type(e).__name__}: {e}")

print("\n完成。")
